"""Build a small energy-twin slide set matching docs/Urban_Regeneration_Final_version.pptx
(Tokyo Smart City Workshop deck): 16:9, Aptos Display titles, GT navy 003057, green accent
3B7A57, bilingual titles. Minimal words, no hyphens, no plus signs.

Slides:
  1. Energy Twin (hero screenshot + 3 points)
  2. Colour by what matters (the 4 metric screenshots, one fixed angle)
  3. How it works (data to experience tech-stack flow)

Run:  python scripts/build_energy_slide.py
Out:  outputs/slides/energy_twin_slide.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / "outputs" / "qa"
OUT = ROOT / "outputs" / "slides" / "energy_twin_slide.pptx"

NAVY = RGBColor(0x00, 0x30, 0x57)
NAVY2 = RGBColor(0x10, 0x2A, 0x43)
INK = RGBColor(0x1E, 0x2A, 0x33)
GREEN = RGBColor(0x3B, 0x7A, 0x57)
GREY = RGBColor(0x6B, 0x76, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xDB, 0xE0)
TF = "Aptos Display"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def box(s, l, t, w, h, anchor=None):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    if anchor: tf.vertical_anchor = anchor
    return tf

def run(p, text, size, color, bold=False, font=TF):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
    r.font.color.rgb = color
    return r

def title(s, en, jp):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(1.16), Inches(2.2), Inches(0.06))
    bar.fill.solid(); bar.fill.fore_color.rgb = GREEN; bar.line.fill.background(); bar.shadow.inherit = False
    tf = box(s, 0.45, 0.26, 12.4, 0.85)
    p = tf.paragraphs[0]
    run(p, en, 30, NAVY, bold=True)
    run(p, "    " + jp, 22, NAVY)

def rrect(s, l, t, w, h, fill, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line: sh.line.color.rgb = line; sh.line.width = Pt(1)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    sh.adjustments[0] = 0.06
    return sh

# ---------------- Slide 1: hero ----------------
s = blank()
title(s, "Energy Twin", "日本橋エネルギーツイン")
s.shapes.add_picture(str(QA / "M0_energy.png"), Inches(0.45), Inches(1.5), width=Inches(8.05))
cap = box(s, 0.45, 6.62, 8.05, 0.3)
run(cap.paragraphs[0], "Nihonbashi energy and carbon model, live demo", 10, GREY)
points = [
    "Every building scored by energy use, intensity, and cost",
    "Test retrofits and cleaner supply, watch the carbon fall",
    "Day and night, on the real Nihonbashi street grid",
]
top = 1.66
for i, txt in enumerate(points, 1):
    n = box(s, 8.8, top, 0.45, 0.6)
    run(n.paragraphs[0], str(i), 22, GREEN, bold=True)
    t = box(s, 9.25, top + 0.02, 3.95, 1.0)
    run(t.paragraphs[0], txt, 15, INK)
    top += 1.15
tk = box(s, 8.8, 5.5, 4.4, 1.2)
run(tk.paragraphs[0], "See where the carbon sits, and prove how to cut it.", 17, NAVY, bold=True)
ft = box(s, 0.45, 7.12, 9, 0.28)
run(ft.paragraphs[0], "Tokyo Smart City Workshop  ·  Nihonbashi 2026", 9, GREY)

# ---------------- Slide 2: colour metrics ----------------
s = blank()
title(s, "Colour by what matters", "指標で色分け")
grid = [
    ("M0_energy.png", "Energy use"),
    ("M1_eui.png", "Energy intensity"),
    ("M2_cost.png", "Revitalisation cost"),
    ("M3_wwr.png", "Window ratio"),
]
IW, IH = 4.3, 2.69
cols = [2.12, 6.92]; rows = [1.62, 4.72]; labrows = [1.35, 4.45]
for k, (img, lab) in enumerate(grid):
    cx = cols[k % 2]; ry = rows[k // 2]; ly = labrows[k // 2]
    s.shapes.add_picture(str(QA / img), Inches(cx), Inches(ry), width=Inches(IW), height=Inches(IH))
    lb = box(s, cx + 0.02, ly, IW, 0.3)
    run(lb.paragraphs[0], lab, 14, NAVY, bold=True)
cp = box(s, 0.45, 7.16, 12, 0.28)
run(cp.paragraphs[0], "Same buildings, same angle. The colour is the only thing that changes.", 11, GREY)

# ---------------- Slide 3: how it works (tech stack flow) ----------------
s = blank()
title(s, "How it works", "可視化の仕組み")
COLS = [
    ("Data", [
        "PLATEAU 3D city, LOD2",
        "OpenStreetMap streets",
        "Building energy model",
        "Supply optimisation, REopt",
    ]),
    ("Processing, Python", [
        "Join and georeference",
        "Convert city tiles to 3D mesh",
        "Build the designed streetscape",
    ]),
    ("3D engine", [
        "CesiumJS georeferenced globe",
        "Three.js WebGL twin",
        "Draco mesh, custom shaders",
    ]),
    ("Experience, browser", [
        "Colour by energy and carbon",
        "Demand and supply scenarios",
        "Day and night, split screen",
        "Live KPIs as you explore",
    ]),
]
cw, gap, x0, hy, by, bh = 2.86, 0.3, 0.5, 1.6, 2.2, 4.5
for i, (head, items) in enumerate(COLS):
    cx = x0 + i * (cw + gap)
    hd = rrect(s, cx, hy, cw, 0.55, NAVY)
    htf = hd.text_frame; htf.word_wrap = True; htf.vertical_anchor = MSO_ANCHOR.MIDDLE
    htf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(htf.paragraphs[0], head, 14, WHITE, bold=True)
    bd = rrect(s, cx, by, cw, bh, WHITE, line=LINE)
    btf = bd.text_frame; btf.word_wrap = True; btf.vertical_anchor = MSO_ANCHOR.TOP
    btf.margin_left = Inches(0.16); btf.margin_right = Inches(0.12); btf.margin_top = Inches(0.16)
    for j, it in enumerate(items):
        p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
        p.space_after = Pt(10)
        run(p, "•  ", 12, GREEN, bold=True)
        run(p, it, 12.5, INK)
    if i < 3:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + cw + 0.02), Inches(by + bh/2 - 0.22), Inches(0.26), Inches(0.44))
        ar.fill.solid(); ar.fill.fore_color.rgb = GREEN; ar.line.fill.background(); ar.shadow.inherit = False
tk = box(s, 0.5, 6.95, 12.3, 0.45)
run(tk.paragraphs[0], "One pipeline, from raw city data to a live model anyone can explore.", 15, NAVY, bold=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"wrote {OUT}  ({OUT.stat().st_size/1000:.0f} KB, {len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
