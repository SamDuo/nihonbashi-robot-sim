"""Build the visual-first slide deck for the Nihonbashi energy twin research.

Source of truth for every number and claim: docs/project_design.md.
Figure provenance: outputs/figures/captions.md.

Studio identity is inherited from scripts/build_energy_slide.py:
GT navy 003057, secondary navy 102A43, green accent 3B7A57, ink 1E2A33,
grey 6B7680, line D6DBE0, Aptos Display, bilingual titles, minimal words.

Run:  python scripts/build_project_slides.py
Out:  outputs/reports/project_design_slides.pptx
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "outputs" / "figures"
PHOTO = FIG / "photoreal"
OUT = ROOT / "outputs" / "reports" / "project_design_slides.pptx"

NAVY = RGBColor(0x00, 0x30, 0x57)
NAVY2 = RGBColor(0x10, 0x2A, 0x43)
INK = RGBColor(0x1E, 0x2A, 0x33)
GREEN = RGBColor(0x3B, 0x7A, 0x57)
GREEN_LT = RGBColor(0x7F, 0xB8, 0x97)
GREY = RGBColor(0x6B, 0x76, 0x80)
LINE = RGBColor(0xD6, 0xDB, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TF = "Aptos Display"

SW, SH = 13.333, 7.5
MARGIN = 0.9            # shared left margin
TITLE_TOP = 0.44        # shared title baseline, identical on every titled slide
RULE_TOP = 1.09         # shared green rule under the title
BODY_TOP = 1.52         # first content row on titled slides

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)


# ---------------------------------------------------------------- primitives
def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def box(s, l, t, w, h, anchor=None, align=None):
    tf = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor:
        tf.vertical_anchor = anchor
    if align:
        tf.paragraphs[0].alignment = align
    return tf


def run(p, text, size, color, bold=False, space=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = TF
    r.font.color.rgb = color
    if space is not None:
        r.font._rPr.set("spc", str(int(space * 100)))
    return r


def rect(s, l, t, w, h, fill=None, line=None, dash=None, lw=1.0):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    style = sh._element.find(qn("p:style"))     # drops the theme effectRef: no drop shadows
    if style is not None:
        sh._element.remove(style)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    sh.shadow.inherit = False
    return sh


def rule(s, l, t, w=2.0, color=GREEN, h=0.05):
    return rect(s, l, t, w, h, fill=color)


def title(s, en, jp=None, size=27, color=NAVY, l=MARGIN, top=TITLE_TOP, rule_color=GREEN):
    tf = box(s, l, top, SW - 2 * l, 0.62, anchor=MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    run(p, en, size, color, bold=True)
    if jp:
        run(p, "    " + jp, size - 8, color)
    rule(s, l, RULE_TOP, 2.0, rule_color)
    return tf


def place_fit(s, path, l, t, w, h, crop_top=0.0):
    """Centre the image inside the box, never stretched. crop_top trims the
    figure's own baked-in title so the slide never carries it twice."""
    iw, ih = Image.open(path).size
    a = iw / (ih * (1 - crop_top))
    fw, fh = (a * h, h) if w / h > a else (w, w / a)
    pic = s.shapes.add_picture(
        str(path), Inches(l + (w - fw) / 2), Inches(t + (h - fh) / 2),
        Inches(fw), Inches(fh))
    if crop_top:
        pic.crop_top = crop_top
    return pic


def place_cover(s, path, l, t, w, h):
    """Fill the box exactly, preserving aspect by cropping the overflow."""
    iw, ih = Image.open(path).size
    a, target = iw / ih, w / h
    pic = s.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))
    if a > target:                      # image too wide: crop the sides
        c = (1 - target / a) / 2
        pic.crop_left = pic.crop_right = c
    elif a < target:                    # image too tall: crop top and bottom
        c = (1 - a / target) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def numeral_slide(numeral, subline, tag, num_pt=180):
    """Full-bleed navy field, one very large number, one line under it."""
    s = blank()
    rect(s, 0, 0, SW, SH, fill=NAVY)
    tf = box(s, MARGIN, 0.72, SW - 2 * MARGIN, 0.4)
    run(tf.paragraphs[0], tag, 13, GREEN_LT, bold=True, space=2.2)
    nb = box(s, MARGIN, 2.05, SW - 2 * MARGIN, 2.9, anchor=MSO_ANCHOR.MIDDLE)
    run(nb.paragraphs[0], numeral, num_pt, WHITE, bold=True, space=-2)
    rule(s, MARGIN, 5.25, 2.0, GREEN, h=0.055)
    sb = box(s, MARGIN, 5.62, SW - 2 * MARGIN - 1.0, 0.9)
    run(sb.paragraphs[0], subline, 19, LINE)
    return s


def chart_slide(headline, path, kicker=None, kicker_big=None, crop_top=0.0):
    """A figure carried at the largest size the slide allows, plus at most one line."""
    s = blank()
    title(s, headline)
    top = 1.40
    bottom = 6.58 if (kicker or kicker_big) else 7.24
    place_fit(s, path, 0.45, top, SW - 0.90, bottom - top, crop_top=crop_top)
    if kicker_big:
        kb = box(s, MARGIN, 6.66, 2.2, 0.68, anchor=MSO_ANCHOR.MIDDLE)
        run(kb.paragraphs[0], kicker_big, 40, GREEN, bold=True, space=-1)
        kt = box(s, MARGIN + 2.3, 6.66, SW - 2 * MARGIN - 2.3, 0.68, anchor=MSO_ANCHOR.MIDDLE)
        run(kt.paragraphs[0], kicker, 17, INK)
    elif kicker:
        kt = box(s, MARGIN, 6.66, SW - 2 * MARGIN, 0.68, anchor=MSO_ANCHOR.MIDDLE)
        run(kt.paragraphs[0], kicker, 17, INK)
    return s


# ------------------------------------------------------------------ 1. title
s = blank()
IMG_H = SW * 836 / 1887                                   # 5.907 in, full width, no crop
rect(s, 0, 0, SW, SH, fill=NAVY2)
s.shapes.add_picture(str(PHOTO / "C1_hero_energy_day.png"), 0, 0, Inches(SW), Inches(IMG_H))
rect(s, 0, IMG_H, SW, SH - IMG_H, fill=NAVY2)             # dark scrim band
rule(s, MARGIN, IMG_H + 0.20, 1.6, GREEN, h=0.05)
jb = box(s, MARGIN, IMG_H + 0.36, SW - 2 * MARGIN, 0.30)
run(jb.paragraphs[0], "日本橋エネルギーツイン", 14, GREEN_LT, space=1.2)
tb = box(s, MARGIN, IMG_H + 0.68, SW - 2 * MARGIN, 0.52)
run(tb.paragraphs[0], "From Georeferenced Digital Twin to Omniverse", 28, WHITE, bold=True)
sb = box(s, MARGIN, IMG_H + 1.22, SW - 2 * MARGIN, 0.30)
run(sb.paragraphs[0], "Sam Duong  ·  Advisor: Dr. Perry Yang", 12.5, LINE)

# --------------------------------------------------------------- 2. question
s = blank()
title(s, "The question", "問い")
qt = box(s, MARGIN, BODY_TOP + 0.22, SW - 2 * MARGIN, 2.1)
qt.paragraphs[0].line_spacing = 1.16
run(qt.paragraphs[0],
    "Does per-building energy use intensity, with ping-derived occupancy, "
    "materially change the cost-optimal PV and battery sizing for a Nihonbashi block?",
    26, NAVY, bold=True)
rect(s, MARGIN, 4.16, SW - 2 * MARGIN, 0.015, fill=LINE)
CRIT = [("10%", "PV capacity change"), ("10%", "storage energy change"),
        ("5", "buildings change rank"), ("2", "top-5 members change")]
colw = (SW - 2 * MARGIN) / 4
for i, (num, lab) in enumerate(CRIT):
    cx = MARGIN + i * colw
    if i:
        rect(s, cx - 0.18, 4.72, 0.012, 1.52, fill=LINE)
    nb = box(s, cx, 4.64, colw - 0.40, 1.20, anchor=MSO_ANCHOR.MIDDLE)
    run(nb.paragraphs[0], num, 62, NAVY, bold=True, space=-1.5)
    lb = box(s, cx, 5.94, colw - 0.40, 0.6)
    run(lb.paragraphs[0], lab, 15, GREY)

# ---------------------------------------------------- 3 & 4. the two numbers
numeral_slide("49.93", "every building, identical intensity  ·  kWh/m²/yr",
              "THE PRO-RATA SPLIT", num_pt=185)
numeral_slide("14 of 22",
              "buildings change rank under corrected intensities  ·  ρ = 0.897 < 0.9",
              "RANK DISPLACEMENT", num_pt=150)

# -------------------------------------------------------------- 5. evolution
s = blank()
title(s, "Context changed, not data")
EV_H = SW * 946 / 3603                                    # 3.50 in, full width, no crop
s.shapes.add_picture(str(PHOTO / "C0_evolution_before_after.png"), 0,
                     Inches(BODY_TOP + (7.16 - BODY_TOP - EV_H) / 2), Inches(SW), Inches(EV_H))

# ------------------------------------------------------------ 6, 7, 8. charts
chart_slide("Building energy use intensity by scenario",
            FIG / "F1_building_energy_intensity.png", crop_top=100 / 1826)
chart_slide("Annual saving under S2 by building",
            FIG / "F2_retrofit_priority_s2.png", crop_top=100 / 1733,
            kicker="Largest absolute savers: 4050, 3042, 3554, 3562, 4571")
chart_slide("District electricity: composition and seasonality",
            FIG / "F3_load_composition_seasonality.png",
            kicker="of the district total is vehicle charging", kicker_big="48.1%")

# ----------------------------------------------------------- 9. design study
s = blank()
GAP = 0.28                                                # images bleed to both edges
IW = (SW - GAP) / 2
IH = IW / 2.28                                            # common aspect, cropped not squashed
top = (SH - IH - 0.66) / 2
for i, (img, lab) in enumerate([
        ("C3_inspector_b4050_s1_energy.png", "Per-building record"),
        ("C6_occupancy_agents_inspector_b3037.png", "Occupancy and agents")]):
    cx = i * (IW + GAP)
    place_cover(s, PHOTO / img, cx, top, IW, IH)
    lb = box(s, cx + 0.55, top + IH + 0.30, IW - 0.55, 0.36)
    run(lb.paragraphs[0], lab, 16, NAVY, bold=True)

# ------------------------------------------------------------ 10. architecture
s = blank()
place_fit(s, FIG / "F7_system_architecture.png", 0.40, 0.40, SW - 0.80, SH - 0.80)

# --------------------------------------------------------- 11. fidelity ladder
s = blank()
title(s, "The fidelity ladder", "忠実度の階層")
ROWS = [("L1", "Dashboard", "built"),
        ("L2", "Web digital twin", "built"),
        ("L3", "Simulation twin", "designed")]
rh, rgap = 0.98, 0.30
ry0 = 2.28
for i, (tier, label, state) in enumerate(reversed(ROWS)):
    ry = ry0 + i * (rh + rgap)
    indent = MARGIN + (len(ROWS) - 1 - i) * 0.62
    w = SW - MARGIN - indent
    built = state == "built"
    if built:
        rect(s, indent, ry, w, rh, fill=None, line=NAVY, lw=1.5)
        rect(s, indent, ry, 0.10, rh, fill=NAVY)
    else:
        rect(s, indent, ry, w, rh, fill=None, line=GREY, lw=1.5, dash=4)
    tb = box(s, indent + 0.40, ry, 0.9, rh, anchor=MSO_ANCHOR.MIDDLE)
    run(tb.paragraphs[0], tier, 22, GREEN if built else GREY, bold=True)
    lb = box(s, indent + 1.35, ry, w - 1.75, rh, anchor=MSO_ANCHOR.MIDDLE)
    run(lb.paragraphs[0], label, 21, NAVY if built else GREY, bold=built)
hb = box(s, MARGIN, 6.52, SW - 2 * MARGIN, 0.5)
run(hb.paragraphs[0],
    "Hypothesis: D-1 settles at L1 or L2, D-2 settles at L2, and only D-3 requires L3.",
    17, INK)

# ------------------------------------------------------------------ 12. close
s = blank()
rect(s, 0, 0, SW, SH, fill=NAVY)
qb = box(s, MARGIN, 2.32, SW - 2 * MARGIN - 0.8, 1.9)
qb.paragraphs[0].line_spacing = 1.14
run(qb.paragraphs[0],
    "At which fidelity tier does the answer to a given urban energy decision stop changing?",
    34, WHITE, bold=True)
rule(s, MARGIN, 4.62, 2.0, GREEN, h=0.055)
cb = box(s, MARGIN, 5.02, SW - 2 * MARGIN, 0.46)
run(cb.paragraphs[0], "CUPUM 2027", 22, GREEN_LT, bold=True, space=1.5)
pb = box(s, MARGIN, 5.62, SW - 2 * MARGIN, 0.4)
run(pb.paragraphs[0],
    "github.com/SamDuo/nihonbashi-robot-sim  ·  demo: outputs/energy_cesium_view.html",
    14, LINE)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"wrote {OUT}  ({OUT.stat().st_size / 1000:.0f} KB, {len(prs.slides._sldIdLst)} slides)")
