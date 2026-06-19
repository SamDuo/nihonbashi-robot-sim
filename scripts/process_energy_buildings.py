"""Process the per-building UBEM workbooks (Nihonbashi_District/*.xlsx) into a
compact per-building energy profile: end-use split, a 24-hour average load shape,
monthly totals, plus the UBEM parameters (WWR, material, R-value) from
Index_baseline.xlsx.

The hourly workbooks appear to be archetype-model outputs (their absolute annual
sums do NOT match the official per-building totals and don't scale with GFA), so
we extract SHAPE (profiles, end-use proportions) — which is unambiguous — and
leave absolute magnitude to energy_dataset.json. Writes data/energy/building_energy.json.
"""
from __future__ import annotations
import json, glob, os
from pathlib import Path
import openpyxl

ED = Path(__file__).resolve().parents[1] / "data" / "energy"
DIR = ED / "Nihonbashi_District"

# ---- UBEM parameters from the index --------------------------------------
idx = {}
wb = openpyxl.load_workbook(DIR / "Index_baseline.xlsx", data_only=True, read_only=True)
ws = wb["Sheet1"]
hdr = [str(c).strip().lower() if c else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
col = {h: i for i, h in enumerate(hdr)}
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[col["id"]] is None:
        continue
    bid = str(int(row[col["id"]]))
    idx[bid] = {
        "material": row[col.get("material", 2)],
        "r_value": row[col.get("r-value", 3)],
        "wwr": row[col.get("window ratio", 4)],
        "use": row[col.get("use", 5)],
        "hvac": row[col.get("hvac", 6)],
        "program": row[col.get("program", 7)],
        "people": row[col.get("people", 9)] if "people" in col else None,
    }
print(f"index: {len(idx)} buildings; WWR range "
      f"{min(b['wwr'] for b in idx.values())}–{max(b['wwr'] for b in idx.values())}")

# ---- per-building hourly workbooks → shape --------------------------------
DAYS_IN_MONTH = [31,28,31,30,31,30,31,31,30,31,30,31]
def month_of_day(d):           # d = 1..365
    m = 0
    while d > DAYS_IN_MONTH[m]:
        d -= DAYS_IN_MONTH[m]; m += 1
    return m

out = {}
for path in sorted(glob.glob(str(DIR / "*.xlsx"))):
    name = os.path.basename(path)
    if name.startswith("Index") or name.startswith("~$"):
        continue
    bid = name[:-5]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    hour24 = [0.0]*24            # avg total by hour-of-day
    enduse = {"cooling":0.0, "heating":0.0, "lighting":0.0}
    monthly = [0.0]*12
    h = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        c, ht, l = float(row[0] or 0), float(row[1] or 0), float(row[2] or 0)
        tot = c + ht + l
        enduse["cooling"] += c; enduse["heating"] += ht; enduse["lighting"] += l
        hour24[h % 24] += tot
        monthly[month_of_day(h // 24 + 1)] += tot
        h += 1
    total = sum(enduse.values()) or 1.0
    mx = max(hour24) or 1.0
    out[bid] = {
        **idx.get(bid, {}),
        "enduse_frac": {k: round(v/total, 4) for k, v in enduse.items()},
        "profile24": [round(x/mx, 4) for x in hour24],          # 0..1 load shape
        "monthly_frac": [round(x/(sum(monthly) or 1), 4) for x in monthly],
        "_archetype_annual_raw": round(total, 1),               # flagged: not building total
    }
    print(f"  {bid}: split C{out[bid]['enduse_frac']['cooling']:.2f}/"
          f"H{out[bid]['enduse_frac']['heating']:.2f}/L{out[bid]['enduse_frac']['lighting']:.2f}"
          f"  WWR={out[bid].get('wwr')}")

json.dump(out, open(ED / "building_energy.json", "w"), separators=(",", ":"), ensure_ascii=False)
print(f"\nwrote {ED/'building_energy.json'}  ({len(out)} buildings)")
